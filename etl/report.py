"""REPORT: build the Excel workbook (live formulas over the cleaned data)."""

from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .rules import OUT_XLSX as OUT, is_blank


# ----------------------------------------------------------------- Excel report
FONT = "Arial"
NAVY = "1F3864"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_sub = Font(name=FONT, size=9, italic=True, color="595959")
f_sec = Font(name=FONT, size=11, bold=True, color=NAVY)
f_kpi_l = Font(name=FONT, size=9, bold=True, color="595959")
f_kpi_v = Font(name=FONT, size=18, bold=True, color=NAVY)
fill_head = PatternFill("solid", fgColor=NAVY)
fill_total = PatternFill("solid", fgColor="D9E1F2")
fill_kpi = PatternFill("solid", fgColor="EEF3FA")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '$#,##0.00;($#,##0.00);-'
MONEY0 = '$#,##0;($#,##0);-'


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = f_head, fill_head, box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_workbook(clean, rej, quality, received, dups, n_sources):
    wb = Workbook()
    n = len(clean)
    last = n + 1

    # ---------------- Clean Data
    cd = wb.active
    cd.title = "Clean Data"
    heads = ["Order ID", "Order Date", "Month", "Customer", "Product", "Channel",
             "Quantity", "Unit Price", "Discount", "Revenue", "Source File"]
    cd.append(heads)
    for i, r in enumerate(clean.itertuples(index=False), start=2):
        cd.append([r[0], r[1], None, r[2], r[3], r[4], r[5], r[6], r[7], None, r[8]])
        cd[f"C{i}"] = f"=DATE(YEAR(B{i}),MONTH(B{i}),1)"
        cd[f"J{i}"] = f"=ROUND(G{i}*H{i}*(1-I{i}),2)"
    style_header(cd, 1, len(heads))
    for row in cd.iter_rows(min_row=2, max_row=last, max_col=len(heads)):
        for cell in row:
            cell.font, cell.border = f_base, box
        row[1].number_format = "yyyy-mm-dd"
        row[2].number_format = "mmm yyyy"
        row[7].number_format = MONEY
        row[8].number_format = "0%"
        row[9].number_format = MONEY
        row[6].alignment = Alignment(horizontal="center")
    for col, w in zip("ABCDEFGHIJK", [13, 13, 11, 22, 30, 12, 11, 13, 11, 14, 18]):
        cd.column_dimensions[col].width = w
    cd.freeze_panes = "A2"
    cd.auto_filter.ref = f"A1:K{last}"

    CD = "'Clean Data'!"
    rev = f"{CD}$J$2:$J${last}"
    mon = f"{CD}$C$2:$C${last}"
    qty = f"{CD}$G$2:$G${last}"
    prd = f"{CD}$E$2:$E${last}"
    chn = f"{CD}$F$2:$F${last}"

    # ---------------- Summary
    s = wb.create_sheet("Summary", 0)
    s.sheet_view.showGridLines = False
    dmin, dmax = clean["Order Date"].min(), clean["Order Date"].max()
    s["A1"] = "Sales Performance Report"
    s["A1"].font = f_title
    s["A2"] = (f"Consolidated from {n_sources} source files  |  {dmin:%d %b %Y} to {dmax:%d %b %Y}"
               "  |  SAMPLE DATA: synthetic, for demonstration")
    s["A2"].font = f_sub

    kpis = [("TOTAL REVENUE", f"=SUM({rev})", MONEY0),
            ("ORDERS", f"=COUNTA({CD}$A$2:$A${last})", "#,##0"),
            ("UNITS SOLD", f"=SUM({qty})", "#,##0"),
            ("AVG ORDER VALUE", "=IF(B5=0,0,A5/B5)", MONEY)]
    for i, (label, formula, fmt) in enumerate(kpis):
        col = "ABCD"[i]
        s[f"{col}4"] = label
        s[f"{col}5"] = formula
        s[f"{col}4"].font, s[f"{col}5"].font = f_kpi_l, f_kpi_v
        s[f"{col}4"].fill = s[f"{col}5"].fill = fill_kpi
        s[f"{col}5"].number_format = fmt
        s[f"{col}4"].alignment = s[f"{col}5"].alignment = Alignment(horizontal="center")
    s.row_dimensions[5].height = 30

    s["A7"] = "Monthly trend"
    s["A7"].font = f_sec
    for c, h in enumerate(["Month", "Revenue ($)", "Orders", "Units", "Avg Order ($)"], start=1):
        s.cell(row=8, column=c, value=h)
    style_header(s, 8, 5)
    months = pd.period_range(dmin, dmax, freq="M")
    r0 = 9
    for i, p in enumerate(months):
        r = r0 + i
        s[f"A{r}"] = date(p.year, p.month, 1)
        s[f"B{r}"] = f"=SUMIFS({rev},{mon},A{r})"
        s[f"C{r}"] = f"=COUNTIFS({mon},A{r})"
        s[f"D{r}"] = f"=SUMIFS({qty},{mon},A{r})"
        s[f"E{r}"] = f"=IF(C{r}=0,0,B{r}/C{r})"
    rl = r0 + len(months) - 1
    tr = rl + 1
    s[f"A{tr}"] = "Total"
    for col in "BCD":
        s[f"{col}{tr}"] = f"=SUM({col}{r0}:{col}{rl})"
    s[f"E{tr}"] = f"=IF(C{tr}=0,0,B{tr}/C{tr})"
    for r in range(r0, tr + 1):
        for c in range(1, 6):
            cell = s.cell(row=r, column=c)
            cell.font = f_bold if r == tr else f_base
            cell.border = box
            if r == tr:
                cell.fill = fill_total
        s[f"A{r}"].number_format = "mmm yyyy"
        s[f"A{r}"].alignment = Alignment(horizontal="left")
        s[f"B{r}"].number_format = MONEY0
        s[f"C{r}"].number_format = s[f"D{r}"].number_format = "#,##0"
        s[f"E{r}"].number_format = MONEY

    # products (sorted by revenue at build time; values are live formulas)
    order = (clean.assign(rev=clean.Quantity * clean["Unit Price"] * (1 - clean.Discount))
             .groupby("Product")["rev"].sum().sort_values(ascending=False).index.tolist())
    pr0 = tr + 3
    s[f"A{pr0 - 1}"] = "Revenue by product"
    s[f"A{pr0 - 1}"].font = f_sec
    for c, h in enumerate(["Product", "Units", "Revenue ($)", "% of Revenue"], start=1):
        s.cell(row=pr0, column=c, value=h)
    style_header(s, pr0, 4)
    for i, name in enumerate(order, start=pr0 + 1):
        s[f"A{i}"] = name
        s[f"B{i}"] = f"=SUMIFS({qty},{prd},A{i})"
        s[f"C{i}"] = f"=SUMIFS({rev},{prd},A{i})"
        s[f"D{i}"] = f"=IF($A$5=0,0,C{i}/$A$5)"
        for c in range(1, 5):
            s.cell(row=i, column=c).font, s.cell(row=i, column=c).border = f_base, box
        s[f"B{i}"].number_format = "#,##0"
        s[f"C{i}"].number_format = MONEY0
        s[f"D{i}"].number_format = "0.0%"
    pl = pr0 + len(order)

    # channels
    ch0 = pl + 3
    s[f"A{ch0 - 1}"] = "Revenue by channel"
    s[f"A{ch0 - 1}"].font = f_sec
    for c, h in enumerate(["Channel", "Orders", "Revenue ($)", "% of Revenue"], start=1):
        s.cell(row=ch0, column=c, value=h)
    style_header(s, ch0, 4)
    for i, name in enumerate(sorted(clean["Channel"].unique()), start=ch0 + 1):
        s[f"A{i}"] = name
        s[f"B{i}"] = f"=COUNTIFS({chn},A{i})"
        s[f"C{i}"] = f"=SUMIFS({rev},{chn},A{i})"
        s[f"D{i}"] = f"=IF($A$5=0,0,C{i}/$A$5)"
        for c in range(1, 5):
            s.cell(row=i, column=c).font, s.cell(row=i, column=c).border = f_base, box
        s[f"B{i}"].number_format = "#,##0"
        s[f"C{i}"].number_format = MONEY0
        s[f"D{i}"].number_format = "0.0%"
    s[f"A{ch0 + 5}"] = ("Revenue = Quantity x Unit Price x (1 - Discount). All figures on this sheet are "
                        "formulas over the Clean Data sheet.")
    s[f"A{ch0 + 5}"].font = f_sub
    for col, w in zip("ABCDE", [30, 16, 16, 16, 16]):
        s.column_dimensions[col].width = w

    def bar(title, cats, vals, colors, anchor, w=15, h=7.5):
        ch = BarChart()
        ch.type = "col"
        ch.title = title
        ch.legend = None
        ch.varyColors = False
        ch.add_data(vals, titles_from_data=False)
        ch.set_categories(cats)
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.majorGridlines = None
        ch.y_axis.number_format = '$#,##0'
        ser = ch.series[0]
        ser.graphicalProperties.solidFill = colors
        ser.graphicalProperties.line.solidFill = colors
        ch.gapWidth = 50
        ch.width, ch.height = w, h
        s.add_chart(ch, anchor)

    bar("Revenue by month", Reference(s, min_col=1, min_row=r0, max_row=rl),
        Reference(s, min_col=2, min_row=r0, max_row=rl), "2F5597", "G4")
    bar("Revenue by product", Reference(s, min_col=1, min_row=pr0 + 1, max_row=pl),
        Reference(s, min_col=3, min_row=pr0 + 1, max_row=pl), "70AD47", "G20", h=9)

    # ---------------- Data Quality Log
    dq = wb.create_sheet("Data Quality Log")
    dq.sheet_view.showGridLines = False
    dq["A1"] = "Data Quality Log"
    dq["A1"].font = f_title
    dq["A2"] = "What the pipeline found and fixed, with a row-count cross-check."
    dq["A2"].font = f_sub
    dq["A4"] = "Row-count cross-check"
    dq["A4"].font = f_sec
    rows = [("Rows received (all source files)", received),
            ("Duplicate rows removed", dups),
            ("Rows rejected (see Rejected Rows sheet)", len(rej)),
            ("Rows in Clean Data", f"=COUNTA('Clean Data'!$A$2:$A${last})"),
            ("Difference (must be 0)", "=B5-B6-B7-B8")]
    for i, (label, val) in enumerate(rows, start=5):
        dq[f"A{i}"] = label
        dq[f"B{i}"] = val
        dq[f"A{i}"].font = f_bold if i == 9 else f_base
        dq[f"B{i}"].font = f_bold if i == 9 else f_base
        dq[f"A{i}"].border = dq[f"B{i}"].border = box
        dq[f"B{i}"].number_format = "#,##0"
    dq["C9"] = '=IF(B9=0,"OK - every row accounted for","CHECK")'
    dq["C9"].font = Font(name=FONT, size=10, bold=True, color="2E7D32")

    dq["A11"] = "Cleaning steps applied"
    dq["A11"].font = f_sec
    for c, h in enumerate(["Check", "Rows / cells affected", "Action taken"], start=1):
        dq.cell(row=12, column=c, value=h)
    style_header(dq, 12, 3)
    for i, (check, cnt, action) in enumerate(quality, start=13):
        dq[f"A{i}"], dq[f"B{i}"], dq[f"C{i}"] = check, cnt, action
        for c in range(1, 4):
            dq.cell(row=i, column=c).font, dq.cell(row=i, column=c).border = f_base, box
        dq[f"B{i}"].number_format = "#,##0"
        dq[f"B{i}"].alignment = Alignment(horizontal="center")
    for col, w in zip("ABC", [46, 22, 62]):
        dq.column_dimensions[col].width = w

    # ---------------- Rejected Rows
    rj = wb.create_sheet("Rejected Rows")
    cols = ["Source", "Order ID", "Order Date (raw)", "Product (raw)", "Quantity (raw)",
            "Unit Price (raw)", "Reason"]
    rj.append(cols)
    for r in rej.itertuples(index=False):
        rj.append([("" if is_blank(v) else str(v)) for v in r])
    style_header(rj, 1, len(cols))
    for row in rj.iter_rows(min_row=2, max_row=len(rej) + 1, max_col=len(cols)):
        for cell in row:
            cell.font, cell.border = f_base, box
    for col, w in zip("ABCDEFG", [20, 13, 17, 30, 15, 16, 32]):
        rj.column_dimensions[col].width = w
    rj.freeze_panes = "A2"
    rj.auto_filter.ref = f"A1:G{len(rej) + 1}"

    wb.save(OUT)


